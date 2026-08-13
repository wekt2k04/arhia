# -*- coding: utf-8 -*-
"""Agent 3bis : Animations & transitions — VERSION 3.

Vision v3 :
  - Morphose (Morph) en continu sur la quasi-totalite du deck (2..26) : des
    continuites fluides d'une abstraction a l'autre.
  - Revelations purement sequentielles (au clic) pour les diagrammes et le code :
    pipeline Agent par Agent, toctou requete par requete, architecture anneau
    par anneau, use case etape par etape, editor/mockup en un bloc.
  - Aucune animation superflue (rebond, spirale...).
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from agent_base import AgentBase

try:
    from pptx import Presentation
    from pptx.oxml import parse_xml
except ImportError:  # pragma: no cover
    Presentation = None

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
NSDECL = (f'xmlns:p="{NS_P}" xmlns:a="{NS_A}" xmlns:r="{NS_R}" xmlns:p14="{NS_P14}"')

DUR_MS = 250

SLIDES_MORPH = set(range(2, 27))  # morph en continu

# Revelations sequentielles : slide -> liste de prefixes de noms de formes
SEQUENCES = {
    2: ["kpi_0", "kpi_1", "kpi_2"],
    8: ["ring_0", "ring_1", "ring_2", "ring_3", "port"],
    9: ["pipe_0", "pipe_1", "pipe_2", "pipe_3", "pipe_4", "sse"],
    10: ["req1", "req2", "idx", "rej", "db"],
    12: ["code"],
    13: ["code"],
    14: ["av_0", "av_1", "av_2", "av_3", "av_4", "cap"],
    15: ["mockup"],
    17: ["kpi_0", "kpi_1", "kpi_2", "kpi_3"],
}


class AnimateurAgent(AgentBase):
    agent_id = "agent_animateur"
    name = "L'Architecte (animations) v3"

    @staticmethod
    def _effect_xml(spid: int, node_id: int, grp_id: int, node_type: str,
                    dur: int = DUR_MS) -> str:
        cond = "indefinite" if node_type == "clickEffect" else "0"
        return f"""<p:par>
  <p:cTn id="{node_id}" fill="hold">
    <p:stCondLst><p:cond delay="{cond}"/></p:stCondLst>
    <p:childTnLst>
      <p:par>
        <p:cTn id="{node_id + 1}" fill="hold">
          <p:stCondLst><p:cond delay="0"/></p:stCondLst>
          <p:childTnLst>
            <p:par>
              <p:cTn id="{node_id + 2}" presetID="10" presetClass="entr"
                     presetSubtype="0" fill="hold" grpId="{grp_id}"
                     nodeType="{node_type}">
                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                <p:childTnLst>
                  <p:set>
                    <p:cBhvr>
                      <p:cTn id="{node_id + 3}" dur="{dur}" fill="hold">
                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                      </p:cTn>
                      <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                      <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                    </p:cBhvr>
                    <p:to><p:strVal val="visible"/></p:to>
                  </p:set>
                  <p:set>
                    <p:cBhvr>
                      <p:cTn id="{node_id + 4}" dur="{dur}" fill="hold">
                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                      </p:cTn>
                      <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
                      <p:attrNameLst><p:attrName>style.opacity</p:attrName></p:attrNameLst>
                    </p:cBhvr>
                    <p:to><p:strVal val="100000"/></p:to>
                  </p:set>
                </p:childTnLst>
              </p:cTn>
            </p:par>
          </p:childTnLst>
        </p:cTn>
      </p:par>
    </p:childTnLst>
  </p:cTn>
</p:par>"""

    def _timing_xml(self, effects: list[dict]) -> str:
        node_id = 10
        groups, bld = [], []
        for eff in effects:
            groups.append(self._effect_xml(eff["spid"], node_id, eff["grp_id"], "clickEffect"))
            node_id += 10
            bld.append(f'<p:bldP spid="{eff["spid"]}" grpId="{eff["grp_id"]}"/>')
        return f"""<p:timing {NSDECL}>
  <p:tnLst>
    <p:par>
      <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
        <p:childTnLst>
          <p:seq concurrent="1" nextAc="seek">
            <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
              <p:childTnLst>
                {''.join(groups)}
              </p:childTnLst>
            </p:cTn>
            <p:prevCondLst>
              <p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:prevCondLst>
            <p:nextCondLst>
              <p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>
            </p:nextCondLst>
          </p:seq>
          <p:bldLst>
            {''.join(bld)}
          </p:bldLst>
        </p:childTnLst>
      </p:cTn>
    </p:par>
  </p:tnLst>
</p:timing>"""

    @staticmethod
    def _morph_xml() -> str:
        return (f'<p:transition {NSDECL} spd="med" p14:dur="700">'
                f'<p14:morph option="byObject"/></p:transition>')

    def _matches(self, shape_name: str, key: str) -> bool:
        return shape_name == key or shape_name.startswith(key)

    def _animate_slide(self, slide, slide_idx: int) -> int:
        keys = SEQUENCES.get(slide_idx)
        if not keys:
            return 0
        effects = []
        used = set()
        grp = 0
        for key in keys:
            matched = [sh for sh in slide.shapes
                       if self._matches(sh.name, key) and sh.shape_id not in used]
            if not matched:
                continue
            for sh in matched:
                used.add(sh.shape_id)
                effects.append({"spid": sh.shape_id, "grp_id": grp})
            grp += 1
        if not effects:
            return 0
        timing_el = parse_xml(self._timing_xml(effects))
        slide._element.append(timing_el)
        if slide_idx in SLIDES_MORPH:
            trans = parse_xml(self._morph_xml())
            slide._element.insert(slide._element.index(timing_el), trans)
        self.info(f"Slide {slide_idx}: {grp} revelations sequentielles + Morph")
        return grp

    def run(self) -> dict:
        self.info(f"=== {self.name} - ETAPE 5 v3 ===")
        src = self.root / self.config["paths"]["draft"]
        if not src.exists():
            return self.make_report("SKIP", {})
        prs = Presentation(str(src))
        total = 0
        for idx, slide in enumerate(prs.slides, start=1):
            if idx in SLIDES_MORPH and idx not in SEQUENCES:
                trans = parse_xml(self._morph_xml())
                slide._element.append(trans)
            total += self._animate_slide(slide, idx)
        target = self.root / self.config["paths"]["animated"]
        prs.save(target)
        static = self.root / "output" / "draft_static.pptx"
        shutil.copy2(src, static)
        self.info(f"Animated v3: {target} | statique: {static}")
        return self.make_report("OK", {
            "animated": str(target), "static": str(static),
            "morph_slides": len(SLIDES_MORPH),
            "sequences": list(SEQUENCES.keys()),
            "reveals": total,
        })
