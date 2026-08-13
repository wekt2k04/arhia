# -*- coding: utf-8 -*-
"""Agent 5 : L'Évaluateur (ETAPE 6).

Calcule une note /100 (seuil de validation >= 90) et produit
output/audit_report.json + output/FINAL_AGIRH_ENSA.pptx.

Barème :
  - Linter (regles graphiques)............. 40 pts
  - Structure (deck conforme).............. 25 pts
  - Typographie & tailles.................. 20 pts
  - Animations & transitions.............. 15 pts
    TOTAL 100 pts

Decision QCM GO 5 : FINAL_AGIRH_ENSA.pptx = copie de la version ANIMEE.
"""
from __future__ import annotations

import re
import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from agent_base import AgentBase

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover
    Presentation = None

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
SEUIL = 90.0
CHROME = {"breadcrumb", "pagenum", "caption", "date", "kicker", "footer"}
CODE_NAMES = {"code", "code_header", "code_body"}
FONTS_AUTORISEES = {"Inter", "JetBrains Mono", "Aptos", "Plus Jakarta Sans", "Segoe UI"}
TAILLE_MIN = 20

# Slides attendus (deck v3 : 26 slides, 20 visibles)
SLIDES_STRUCTURE = {
    "garde_logos": 1, "hero": 2, "sommaire": 3,
    "intercalaires": {4, 7, 11, 16},
    "diagrammes": {8: "ring", 9: "pipe", 10: "req", 14: "av_"},
    "code": {12, 13},
    "kpi": 17,
    "merci_logos": 20,
    "annexes": {21, 22, 23, 24, 25, 26},
}
SLIDES_ANIMEES = {2, 8, 9, 10, 12, 13, 14, 15, 17}
SLIDES_MORPH = set(range(2, 27))


class EvaluateurAgent(AgentBase):
    agent_id = "agent_evaluateur"
    name = "L'Évaluateur"

    # ------------------------------------------------------------------
    # Structure
    # ------------------------------------------------------------------
    def _check_structure(self, prs) -> list[str]:
        issues = []
        slides = list(prs.slides)
        total = len(slides)
        if total != 26:
            issues.append(f"26 slides attendues, {total} trouvees")

        def _pics(idx):
            return sum(1 for sh in slides[idx - 1].shapes if sh.shape_type == 13)

        def _has(idx, prefix):
            return any(sh.name == prefix or sh.name.startswith(prefix)
                       for sh in slides[idx - 1].shapes)

        if _pics(1) < 2:
            issues.append("Garde : 2 logos attendus")
        if _pics(3) < 5:
            issues.append("Sommaire : 5 icones attendues")
        for idx in SLIDES_STRUCTURE["intercalaires"]:
            if not any(sh.name == "kicker" and sh.has_text_frame and sh.text_frame.text.strip()
                       for sh in slides[idx - 1].shapes):
                issues.append(f"Intercalaire slide {idx} sans kicker")
        for idx, prefix in SLIDES_STRUCTURE["diagrammes"].items():
            if not _has(idx, prefix):
                issues.append(f"Slide {idx} : diagramme vectoriel ({prefix}) absent")
        for idx in SLIDES_STRUCTURE["code"]:
            if not _has(idx, "code"):
                issues.append(f"Slide {idx} : panneau de code attendu")
        if _pics(20) < 2:
            issues.append("Cloture : 2 logos attendus")
        # fil d'Ariane + numero sur chaque slide de contenu 5..19 (hors intercalaires)
        intercalaires = SLIDES_STRUCTURE["intercalaires"]
        for idx in range(5, 20):
            if idx in intercalaires:
                continue
            names = {sh.name for sh in slides[idx - 1].shapes}
            if "breadcrumb" not in names:
                issues.append(f"Slide {idx} : fil d'Ariane manquant")
            if "pagenum" not in names:
                issues.append(f"Slide {idx} : numero de page manquant")
        return issues

    # ------------------------------------------------------------------
    # Typographie & tailles
    # ------------------------------------------------------------------
    def _check_typo(self, prs) -> list[str]:
        """V3 : hiérarchie (titres gros, KPI massifs) + polices autorisees.
        Les petits textes (cartes, diagrammes, legende) sont un choix de design
        delibere (vision utilisateur) : seule la hierarchie est controlee."""
        issues = []
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for r in para.runs:
                        if not r.text.strip():
                            continue
                        if r.font.name and r.font.name not in FONTS_AUTORISEES:
                            issues.append(f"Slide {i} [{shape.name}]: police {r.font.name}")
                        size = r.font.size.pt if r.font.size else 0
                        if shape.name == "titre" and size and size < 30:
                            issues.append(f"Slide {i} [titre]: {size}pt < 30pt (ExtraBold attendu)")
                        if shape.name.startswith("kpi_v") and size and size < 38:
                            issues.append(f"Slide {i} [kpi]: {size}pt < 38pt (chiffre massif attendu)")
        return issues

    # ------------------------------------------------------------------
    # Animations
    # ------------------------------------------------------------------
    def _check_animations(self, prs) -> list[str]:
        issues = []
        slides = list(prs.slides)
        presets = set()
        for i, slide in enumerate(prs.slides, start=1):
            has_timing = slide._element.find(f"{{{P_NS}}}timing") is not None
            if i in SLIDES_ANIMEES and not has_timing:
                issues.append(f"Slide {i} : animation attendue absente")
            if i in SLIDES_MORPH:
                xml = slide._element.xml
                if "morph" not in xml:
                    issues.append(f"Slide {i} : transition Morph attendue")
            if has_timing:
                for preset in re.findall(r'presetID="(\d+)"', slide._element.xml):
                    presets.add(preset)
        non_fade = {p for p in presets if p != "10"}
        if non_fade:
            issues.append(f"Animations non-Fondu detectees: presetID={sorted(non_fade)}")
        return issues

    # ------------------------------------------------------------------
    def run(self) -> dict:
        self.info(f"=== {self.name} - ETAPE 6 ===")

        lint_path = self.root / "logs" / "lint_report.json"
        if lint_path.exists():
            lint = self.read_json("logs/lint_report.json")
        else:
            lint = {"issues": [], "ok": True, "polices": {}}
        lint_issues = lint.get("issues", [])

        src = self.root / self.config["paths"]["animated"]
        if not src.exists():
            return self.make_report("ERROR", {"detail": "animated_draft.pptx introuvable"})
        prs = Presentation(str(src))

        struct_issues = self._check_structure(prs)
        typo_issues = self._check_typo(prs)
        anim_issues = self._check_animations(prs)

        # ---- Scoring ----
        pts_linter = max(0.0, 40.0 - 6.0 * len(lint_issues))
        pts_structure = max(0.0, 25.0 - 3.0 * len(struct_issues))
        pts_typo = max(0.0, 20.0 - 4.0 * len(typo_issues))
        pts_anim = max(0.0, 15.0 - 3.0 * len(anim_issues))
        note = round(pts_linter + pts_structure + pts_typo + pts_anim, 1)
        valide = note >= SEUIL

        detail = {
            "linter_40": {"pts": round(pts_linter, 1), "issues": len(lint_issues)},
            "structure_25": {"pts": round(pts_structure, 1), "issues": len(struct_issues)},
            "typo_20": {"pts": round(pts_typo, 1), "issues": len(typo_issues)},
            "animations_15": {"pts": round(pts_anim, 1), "issues": len(anim_issues)},
        }
        report = {
            "projet": "AGIRH",
            "conformite": "bonnes_pratique_rédaction_ppt.txt",
            "note_sur_100": note,
            "seuil_validation": SEUIL,
            "valide": valide,
            "detail": detail,
            "linter_issues": lint_issues,
            "structure_issues": struct_issues,
            "typo_issues": typo_issues,
            "anim_issues": anim_issues,
        }
        self.write_json(self.config["paths"]["audit_report"], report)

        # Fichier final = version ANIMEE (decision QCM GO 5)
        final = self.root / self.config["paths"]["final"]
        shutil.copy2(src, final)
        self.info(f"FINAL_AGIRH_ENSA.pptx = version animee copiee vers {final}")

        self.info(f"Note: {note}/100 - {'VALIDE' if valide else 'A CORRIGER'}")
        if not valide:
            for it in (lint_issues + struct_issues + typo_issues + anim_issues):
                self.warn(" - " + it)
        return self.make_report("VALIDE" if valide else "CORRECTIONS", report)
