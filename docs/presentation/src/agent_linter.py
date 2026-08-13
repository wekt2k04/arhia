# -*- coding: utf-8 -*-
"""Agent 4 : Le Linter (ETAPE 6).

Parses output/animated_draft.pptx et verifie les regles graphiques actives
(decision QCM GO 5) :
  1. Contraste WCAG >= 4.5:1 (texte de contenu) - regle 27.
  2. Max 3 couleurs dominantes hors neutres (code et chrome exemptes) - regle 22.
  3. Aucun texte bord-a-bord (gouttiere 5%) - regle 33.
  4. Aucun chevauchement involontaire de formes portant du texte - regle 37.

Informations supplementaires (non penalisantes) : polices utilisees.

Sortie : logs/lint_report.json
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from agent_base import AgentBase

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:  # pragma: no cover
    Presentation = None

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

SLIDE_W, SLIDE_H = 13.333, 7.5
GUT, GUT_TB = 0.75, 0.6
MIN_RATIO = 4.5

FOND = (0x0B, 0x13, 0x2B)          # BG #0B132B (contenu)
FOND_DEEP = (0x0A, 0x0F, 0x1C)     # BG_DEEP
SURFACE = (0x11, 0x1B, 0x33)
CODE_BG = (0x0D, 0x11, 0x17)
CODE_HEADER = (0x21, 0x25, 0x2B)

CHROME = {"breadcrumb", "pagenum", "caption", "date", "kicker", "footer"}
CODE_NAMES = {"code", "code_header", "code_body"}
NON_CONTENT = {"decor", "fond"} | CODE_NAMES
SURFACE_NAMES = {"carte", "kpi"}


def rel_luminance(r: int, g: int, b: int) -> float:
    def _lin(c: float) -> float:
        c /= 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def contrast_ratio(rgb1: tuple[int, int, int], rgb2: tuple[int, int, int]) -> float:
    l1, l2 = rel_luminance(*rgb1), rel_luminance(*rgb2)
    if l1 < l2:
        l1, l2 = l2, l1
    return (l1 + 0.05) / (l2 + 0.05)


def is_neutral(rgb: tuple[int, int, int]) -> bool:
    return max(rgb) - min(rgb) < 25


class LinterAgent(AgentBase):
    agent_id = "agent_linter"
    name = "Le Linter"

    def _bg_for(self, shape_name: str) -> tuple[int, int, int]:
        if shape_name == "code_body":
            return CODE_BG
        if shape_name == "code_header":
            return CODE_HEADER
        if any(shape_name.startswith(p) for p in SURFACE_NAMES):
            return SURFACE
        return FOND

    def _collect_runs(self, prs):
        """Retourne [(slide_idx, shape_name, text, rgb, size, shape_obj)]."""
        runs = []
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for r in para.runs:
                        if not r.text.strip():
                            continue
                        try:
                            rgb = tuple(r.font.color.rgb)
                        except Exception:
                            rgb = None
                        size = r.font.size.pt if r.font.size else None
                        runs.append((i, shape.name, r.text, rgb, size, shape))
        return runs

    # ------------------------------------------------------------------
    # Contraste (regle 27) - chrome exempt
    # ------------------------------------------------------------------
    def _check_contrast(self, runs) -> list[str]:
        issues = []
        for i, name, text, rgb, size, _ in runs:
            if name in CHROME or name in NON_CONTENT or rgb is None:
                continue
            ratio = contrast_ratio(rgb, self._bg_for(name))
            if ratio < MIN_RATIO:
                issues.append(
                    f"Slide {i} [{name}]: contraste {ratio:.2f}:1 < {MIN_RATIO}:1 "
                    f"(couleur #{''.join(f'{c:02X}' for c in rgb)}) '{text[:40]}'"
                )
        return issues

    # ------------------------------------------------------------------
    # Max 3 couleurs (regle 22) - neutres, code et chrome exemptes
    # ------------------------------------------------------------------
    def _check_colors(self, runs) -> list[str]:
        colors = set()
        for i, name, text, rgb, size, _ in runs:
            if name in CHROME or name in CODE_NAMES or name in NON_CONTENT or rgb is None:
                continue
            if is_neutral(rgb):
                continue
            colors.add(rgb)
        if len(colors) > 3:
            hexs = ", ".join(f"#{''.join(f'{c:02X}' for c in c)}" for c in sorted(colors))
            return [f"{len(colors)} couleurs dominantes (max 3): {hexs}"]
        return []

    # ------------------------------------------------------------------
    # Texte bord-a-bord (regle 33) - gouttiere 5%
    # ------------------------------------------------------------------
    def _check_borders(self, prs) -> list[str]:
        issues = []
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.name in NON_CONTENT or shape.name in CHROME:
                    continue
                if not shape.has_text_frame or not shape.text_frame.text.strip():
                    continue
                x1 = shape.left / 914400
                y1 = shape.top / 914400
                x2 = (shape.left + shape.width) / 914400
                y2 = (shape.top + shape.height) / 914400
                if (x1 < GUT - 0.05 or y1 < GUT_TB - 0.05 or
                        x2 > SLIDE_W - GUT + 0.05 or y2 > SLIDE_H - GUT_TB + 0.05):
                    issues.append(
                        f"Slide {i} [{shape.name}]: texte hors gouttiere "
                        f"({x1:.2f},{y1:.2f})->({x2:.2f},{y2:.2f})"
                    )
        return issues

    # ------------------------------------------------------------------
    # Chevauchement involontaire (regle 37)
    # ------------------------------------------------------------------
    @staticmethod
    def _overlap(a, b) -> float:
        ax1, ay1, ax2, ay2 = a
        bx1, by1, bx2, by2 = b
        ox = max(0.0, min(ax2, bx2) - max(ax1, bx1))
        oy = max(0.0, min(ay2, by2) - max(ay1, by1))
        return ox * oy

    def _check_overlap(self, prs) -> list[str]:
        issues = []
        for i, slide in enumerate(prs.slides, start=1):
            boxes = []
            for shape in slide.shapes:
                if shape.name in NON_CONTENT or shape.name in CHROME:
                    continue
                if not shape.has_text_frame or not shape.text_frame.text.strip():
                    continue
                boxes.append((shape.name,
                              (shape.left / 914400, shape.top / 914400,
                               (shape.left + shape.width) / 914400,
                               (shape.top + shape.height) / 914400)))
            for a in range(len(boxes)):
                for b in range(a + 1, len(boxes)):
                    na, ba = boxes[a]
                    nb, bb = boxes[b]
                    inter = self._overlap(ba, bb)
                    if inter <= 0:
                        continue
                    area_a = (ba[2] - ba[0]) * (ba[3] - ba[1])
                    area_b = (bb[2] - bb[0]) * (bb[3] - bb[1])
                    if inter > 0.10 * min(area_a, area_b):
                        issues.append(f"Slide {i}: chevauchement {na} / {nb} "
                                      f"({inter / min(area_a, area_b) * 100:.0f}%)")
        return issues

    # ------------------------------------------------------------------
    def _fonts_report(self, prs) -> dict:
        fonts: dict[str, int] = {}
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for r in para.runs:
                        if r.font.name:
                            fonts[r.font.name] = fonts.get(r.font.name, 0) + 1
        return fonts

    # ------------------------------------------------------------------
    def run(self) -> dict:
        self.info(f"=== {self.name} - ETAPE 6 ===")
        self.ensure_rule(27, "contraste WCAG >= 4.5:1 (texte de contenu)")
        self.ensure_rule(22, "max 3 couleurs dominantes (hors neutres)")
        self.ensure_rule(33, "aucun texte bord-a-bord")
        self.ensure_rule(37, "aucun chevauchement involontaire")

        src = self.root / self.config["paths"]["animated"]
        if not src.exists():
            self.warn(f"Animated introuvable: {src}")
            return self.make_report("SKIP", {})

        prs = Presentation(str(src))
        runs = self._collect_runs(prs)
        issues = (
            self._check_contrast(runs)
            + self._check_colors(runs)
            + self._check_borders(prs)
            + self._check_overlap(prs)
        )
        report = {
            "issues": issues,
            "ok": not issues,
            "nb_runs_texte": len(runs),
            "polices": self._fonts_report(prs),
            "regles_actives": ["contrast", "colors", "borders", "overlap"],
        }
        out = self.write_json("logs/lint_report.json", report)
        self.info(f"Lint: {'OK - 0 probleme' if not issues else f'{len(issues)} probleme(s)'}")
        for it in issues:
            self.warn(" - " + it)
        return self.make_report("OK" if not issues else "ISSUES", report)
