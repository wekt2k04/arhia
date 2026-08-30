# -*- coding: utf-8 -*-
"""
Genere docs/presentations/arhia_Soutenance.pptx (soutenance de stage, 20 slides).

Le .pptx genere n'est PAS commite (binaire non-diffable, voir .gitignore) - ce script
source, texte et versionnable, est la reference reproductible.

Dependance externe (hors depot) : les logos AGIRH (entreprise d'accueil) / ENSA Safi
vivent dans C:\\Users\\Wilfried\\OneDrive\\Bureau\\presentation\\assets\\ (fournis par le
porteur du projet, pas versionnes ici - remplacer ASSETS ci-dessous si ce chemin
change). Le logo du produit arhia, lui, vient du depot (frontend/public/arhia-logo.png,
deja versionne) - aucune dependance externe pour celui-la.

Usage : python docs/presentations/generate_pptx.py   (necessite `pip install python-pptx`)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------- Palette (extraite de frontend/app/globals.css, mode sombre reel du produit,
# retravaille le 2026-08-30 - ardoise bleutee + vraie hierarchie carte/fond) ----------
BG          = RGBColor(0x10, 0x12, 0x19)   # --background dark
BG_ALT      = RGBColor(0x1B, 0x1E, 0x27)   # --card dark, nettement plus clair (vraie hierarchie)
PRIMARY     = RGBColor(0x8A, 0x84, 0xF5)   # --primary dark (indigo clair)
PRIMARY_DK  = RGBColor(0x50, 0x48, 0xE5)   # --primary light mode (indigo plus sature)
TEXT        = RGBColor(0xED, 0xEF, 0xF2)   # --foreground dark
MUTED       = RGBColor(0xA5, 0xAB, 0xB6)   # --muted-foreground dark
BORDER      = RGBColor(0x30, 0x34, 0x41)   # --border dark
ACCENT_BG   = RGBColor(0x28, 0x26, 0x4F)   # --accent dark (fond de badge/encadre)
GOOD        = RGBColor(0x39, 0xC6, 0x6D)   # --status-done dark (vert, recalcule WCAG AA)
WARN        = RGBColor(0xFA, 0xB9, 0x47)   # --status-pending dark (orange, recalcule WCAG AA)

FONT = "Segoe UI"
ASSETS = r"C:\Users\Wilfried\OneDrive\Bureau\presentation\assets"
AGIRH_LOGO = ASSETS + r"\agirh_white_rgba.png"
ENSA_LOGO  = ASSETS + r"\ensa_white_rgba.png"
ARHIA_LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "frontend", "public", "arhia-logo.png")

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height

def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # pousser le fond tout en bas de la pile z-order
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s

def set_text(tf, text, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font=FONT, line_spacing=1.0):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing != 1.0:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = font
    return p

def add_textbox(slide, l, t, w, h, text, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font=FONT, anchor=None, line_spacing=1.0):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    set_text(tf, text, size, color, bold, align, font, line_spacing)
    return box

def add_bullets(slide, l, t, w, h, items, size=16, color=TEXT, font=FONT, space_after=10, bullet_color=PRIMARY, line_spacing=1.12):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    first = True
    for item in items:
        if isinstance(item, tuple):
            txt, sub = item
        else:
            txt, sub = item, False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        indent = "      " if sub else ""
        bullet = "–  " if sub else "▪  "
        r = p.add_run()
        r.text = indent + bullet
        r.font.size = Pt(size - (2 if sub else 0)); r.font.color.rgb = bullet_color if not sub else MUTED; r.font.name = font; r.font.bold = not sub
        r2 = p.add_run()
        r2.text = txt
        r2.font.size = Pt(size - (2 if sub else 0)); r2.font.color.rgb = color if not sub else MUTED; r2.font.name = font
    return box

def add_kicker(slide, text, color=PRIMARY):
    add_textbox(slide, Inches(0.7), Inches(0.42), Inches(8), Inches(0.4), text.upper(),
                size=13, color=color, bold=True, font=FONT)

def add_title(slide, text, y=Inches(0.75), size=27, color=TEXT):
    add_textbox(slide, Inches(0.7), y, Inches(11.9), Inches(0.95), text, size=size, color=color, bold=True, font=FONT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), y + Inches(1.05), Inches(1.1), Pt(3))
    line.fill.solid(); line.fill.fore_color.rgb = PRIMARY; line.line.fill.background(); line.shadow.inherit = False
    return line

def add_footer(slide, n, section=""):
    add_textbox(slide, Inches(0.7), Inches(7.08), Inches(6), Inches(0.3), "arhia — Assistant RH agentique" + ((" · " + section) if section else ""),
                size=9, color=MUTED, font=FONT)
    add_textbox(slide, Inches(12.5), Inches(7.08), Inches(0.5), Inches(0.3), str(n), size=9, color=MUTED, font=FONT, align=PP_ALIGN.RIGHT)

def add_logo(slide, path, l, t, h=None, w=None):
    slide.shapes.add_picture(path, l, t, height=h, width=w)

def add_box(slide, l, t, w, h, text, fill=ACCENT_BG, text_color=TEXT, line_color=PRIMARY, size=13, bold=True, align=PP_ALIGN.CENTER, line_w=1.25, corner=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if corner else MSO_SHAPE.RECTANGLE
    box = slide.shapes.add_shape(shape_type, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = fill
    box.line.color.rgb = line_color; box.line.width = Pt(line_w)
    box.shadow.inherit = False
    if corner:
        try:
            box.adjustments[0] = 0.08
        except Exception:
            pass
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(6); tf.margin_right = Pt(6); tf.margin_top = Pt(4); tf.margin_bottom = Pt(4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, text, size=size, color=text_color, bold=bold, align=align)
    return box

def add_arrow_down(slide, cx, t, length=Inches(0.28), color=PRIMARY, w=Pt(2.25)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, cx, t, cx, t + length)
    conn.line.color.rgb = color; conn.line.width = w
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn

def add_arrow_right(slide, l, cy, length=Inches(0.3), color=PRIMARY, w=Pt(2.25)):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, l, cy, l + length, cy)
    conn.line.color.rgb = color; conn.line.width = w
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
    ln.append(tail)
    return conn

def add_badge(slide, l, t, text, fill=PRIMARY, color=RGBColor(0xFF,0xFF,0xFF), size=11, w=Inches(1.5), h=Inches(0.36)):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.fill.background(); b.shadow.inherit = False
    try:
        b.adjustments[0] = 0.5
    except Exception:
        pass
    tf = b.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top = 0; tf.margin_bottom = 0
    set_text(tf, text, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)
    return b

SECTION_NAMES = {1: "Contexte & besoin", 2: "Architecture", 3: "Coeur IA — RAG & orchestration", 4: "Resultats & deploiement"}

# ============================================================ SLIDE 1 — TITRE
s = add_slide()
# bande verticale d'accent a gauche
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), SH)
band.fill.solid(); band.fill.fore_color.rgb = PRIMARY; band.line.fill.background(); band.shadow.inherit = False
add_logo(s, AGIRH_LOGO, Inches(0.9), Inches(0.7), h=Inches(0.62))
add_logo(s, ENSA_LOGO, Inches(10.13), Inches(0.86), w=Inches(2.6))
add_textbox(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(0.5), "PROJET DE FIN D'ANNEE — SOUTENANCE DE STAGE",
            size=14, color=PRIMARY, bold=True)
# Le mot-marque "arhia" vit deja dans le logo (hexagone + texte, degrade indigo) - pas de
# repeter le mot en texte brut a cote, seulement le descriptif qui le complete.
add_logo(s, ARHIA_LOGO, Inches(0.85), Inches(3.05), h=Inches(1.0))
add_textbox(s, Inches(4.55), Inches(3.05), Inches(7.5), Inches(1.0), "Assistant RH Agentique",
            size=32, color=TEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
add_textbox(s, Inches(0.9), Inches(4.05), Inches(11.5), Inches(0.7),
            "Onboarding & Offboarding assistes par IA — pipeline RAG, orchestration conversationnelle et garde-fous anti-hallucination",
            size=16, color=MUTED)
add_textbox(s, Inches(0.9), Inches(6.35), Inches(6), Inches(0.35), "Wilfried TSETSE", size=15, color=TEXT, bold=True)
add_textbox(s, Inches(0.9), Inches(6.68), Inches(7), Inches(0.35),
            "2e annee Genie Informatique de Donnees et IA — ENSA Safi  ·  Encadrant entreprise : M. Moulay Rachid Didi Alaoui", size=11.5, color=MUTED)

# ============================================================ SLIDE 2 — AGENDA
s = add_slide()
add_kicker(s, "Sommaire")
add_title(s, "Plan de la presentation")
items = [
    ("1", "Contexte & besoin metier", "Le probleme, le perimetre, les roles et la vision de l'agent"),
    ("2", "Architecture technique", "Hexagonale, stack, securite (RBAC, JWT, BFF)"),
    ("3", "Coeur IA — RAG & orchestration", "Les 4 phases du pipeline, Router/Generator, garde-fous"),
    ("4", "Resultats, deploiement & perspectives", "Ce qui est verifie, les limites connues, la suite"),
]
y = Inches(2.05)
for num, title, desc in items:
    add_box(s, Inches(0.9), y, Inches(0.62), Inches(0.62), num, fill=PRIMARY, text_color=RGBColor(0xFF,0xFF,0xFF), size=20, corner=True)
    add_textbox(s, Inches(1.75), y - Inches(0.03), Inches(9.8), Inches(0.4), title, size=19, color=TEXT, bold=True)
    add_textbox(s, Inches(1.75), y + Inches(0.38), Inches(9.8), Inches(0.4), desc, size=13, color=MUTED)
    y += Inches(1.18)
add_footer(s, 2)

def section_divider(n, num, title, subtitle):
    s = add_slide()
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    band.fill.solid(); band.fill.fore_color.rgb = BG_ALT; band.line.fill.background(); band.shadow.inherit = False
    sp = band._element; sp.getparent().remove(sp); s.shapes._spTree.insert(3, sp)
    add_textbox(s, Inches(0.9), Inches(2.55), Inches(3), Inches(1.2), f"0{num}", size=90, color=PRIMARY, bold=True)
    add_textbox(s, Inches(3.3), Inches(3.05), Inches(9), Inches(0.9), title, size=34, color=TEXT, bold=True)
    add_textbox(s, Inches(3.32), Inches(3.85), Inches(9), Inches(0.6), subtitle, size=15, color=MUTED)
    add_footer(s, n, SECTION_NAMES[num])
    return s

# ============================================================ SLIDE 3 — SEPARATEUR PARTIE 1
section_divider(3, 1, "Contexte & besoin metier", "Pourquoi un agent IA pour l'onboarding et l'offboarding")

# ============================================================ SLIDE 4 — LE PROBLEME
s = add_slide()
add_kicker(s, "Partie 1 — Contexte")
add_title(s, "Le constat : deux phases critiques, aujourd'hui lentes")
add_bullets(s, Inches(0.7), Inches(2.05), Inches(6.7), Inches(4.2), [
    "L'integration et le depart d'un collaborateur mobilisent 3 acteurs : RH, IT, management",
    "Processus aujourd'hui manuel, source d'oublis (acces IT non revoques, documents non signes...)",
    "Contexte reel : exigence de conformite qualite (checklists SMSI de l'entreprise)",
    "Objectif : un assistant conversationnel qui guide et informe — jamais qui remplace le jugement RH",
], size=16, space_after=16)
box = add_box(s, Inches(8.0), Inches(2.05), Inches(4.6), Inches(4.2), "", fill=ACCENT_BG, line_color=BORDER, corner=True)
add_textbox(s, Inches(8.35), Inches(2.35), Inches(4.0), Inches(0.4), "Ce que l'agent fait", size=14, color=PRIMARY, bold=True)
add_bullets(s, Inches(8.35), Inches(2.85), Inches(3.95), Inches(1.2), [
    "Repondre aux questions sur les politiques internes (documentaire)",
    "Renseigner sur l'avancement d'un dossier personnel",
], size=13, space_after=10)
add_textbox(s, Inches(8.35), Inches(4.35), Inches(4.0), Inches(0.4), "Ce que l'agent ne fait jamais", size=14, color=WARN, bold=True)
add_bullets(s, Inches(8.35), Inches(4.85), Inches(3.95), Inches(1.2), [
    "Declencher une action destructrice ou une ecriture",
    "Inventer une information absente du corpus",
], size=13, space_after=10, bullet_color=WARN)
add_footer(s, 4, SECTION_NAMES[1])

# ============================================================ SLIDE 5 — PERIMETRE & ROLES
s = add_slide()
add_kicker(s, "Partie 1 — Contexte")
add_title(s, "Perimetre resserre : Onboarding / Offboarding, 3 roles")
add_textbox(s, Inches(0.7), Inches(2.0), Inches(11.9), Inches(0.5),
            "Perimetre volontairement strict — pas de conges/CET/paie, pour livrer un socle solide plutot qu'un large prototype fragile.",
            size=13.5, color=MUTED)
roles = [
    ("Collaborateur", "Voit uniquement ses propres donnees (son dossier, sa checklist)"),
    ("RH", "Gere les collaborateurs de son pole uniquement — verifie avant meme le role"),
    ("Admin / Qualite", "Portee globale — seul role habilite a elever un compte ou valider un template"),
]
x = Inches(0.7); w = Inches(3.95)
for i, (role, desc) in enumerate(roles):
    l = x + i * (w + Inches(0.18))
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(2.75), w, Inches(1.55))
    box.fill.solid(); box.fill.fore_color.rgb = ACCENT_BG
    box.line.color.rgb = PRIMARY; box.line.width = Pt(1.25); box.shadow.inherit = False
    try:
        box.adjustments[0] = 0.08
    except Exception:
        pass
    box.text_frame.paragraphs[0].text = ""
    add_textbox(s, l + Inches(0.2), Inches(2.9), w - Inches(0.4), Inches(0.4), role, size=17, color=PRIMARY, bold=True)
    add_textbox(s, l + Inches(0.2), Inches(3.4), w - Inches(0.4), Inches(0.85), desc, size=12, color=MUTED, line_spacing=1.15)
add_textbox(s, Inches(0.7), Inches(4.75), Inches(11.9), Inches(0.4), "Circuit qualite de validation d'un template (checklist)", size=15, color=TEXT, bold=True)
steps = ["Redacteur\n(RH)\npropose", "Verificateur\n(Admin/Qualite)\nvalide", "Approbateur\n(Admin/Qualite)\napprouve"]
x2 = Inches(0.7); w2 = Inches(3.6)
for i, st in enumerate(steps):
    l = x2 + i * (w2 + Inches(0.65))
    add_box(s, l, Inches(5.3), w2, Inches(1.05), st, fill=BG_ALT, line_color=PRIMARY, size=13, corner=True)
    if i < 2:
        add_arrow_right(s, l + w2, Inches(5.3) + Inches(0.52), length=Inches(0.6))
add_textbox(s, Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.4),
            "Seul un template APPROUVE peut instancier un dossier reel — contrainte verifiee par le code, pas seulement documentee.",
            size=12, color=MUTED)
add_footer(s, 5, SECTION_NAMES[1])

# ============================================================ SLIDE 6 — VISION
s = add_slide()
add_kicker(s, "Partie 1 — Contexte")
add_title(s, "La vision : un agent strictement informatif")
add_textbox(s, Inches(0.7), Inches(2.0), Inches(11.9), Inches(0.5),
            "A chaque question, un routeur decide du chemin — jamais le meme traitement pour deux natures de questions differentes.",
            size=13.5, color=MUTED)
add_box(s, Inches(5.15), Inches(2.7), Inches(3.0), Inches(0.85), "Question utilisateur", fill=PRIMARY, text_color=RGBColor(0xFF,0xFF,0xFF), size=15, corner=True)
add_arrow_down(s, Inches(6.65), Inches(3.55), length=Inches(0.35))
add_box(s, Inches(5.0), Inches(3.9), Inches(3.3), Inches(0.85), "Router\n(classification d'intention)", fill=ACCENT_BG, line_color=PRIMARY, size=14, corner=True)
add_arrow_down(s, Inches(3.0), Inches(4.9), length=Inches(0.0))
# branches
conn1 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.6), Inches(4.75), Inches(2.6), Inches(5.15))
conn1.line.color.rgb = PRIMARY; conn1.line.width = Pt(2.25)
conn2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.7), Inches(4.75), Inches(10.6), Inches(5.15))
conn2.line.color.rgb = PRIMARY; conn2.line.width = Pt(2.25)
add_box(s, Inches(0.9), Inches(5.15), Inches(3.4), Inches(1.15), "DOCUMENTAIRE\n→ pipeline RAG (Partie 3)", fill=BG_ALT, line_color=GOOD, size=13, corner=True)
add_box(s, Inches(5.05), Inches(5.15), Inches(3.4), Inches(1.15), "STATUT_DOSSIER\n→ lecture seule d'un dossier", fill=BG_ALT, line_color=GOOD, size=13, corner=True)
add_box(s, Inches(9.0), Inches(5.15), Inches(3.4), Inches(1.15), "HORS_PERIMETRE\n→ refus poli", fill=BG_ALT, line_color=WARN, size=13, corner=True)
add_textbox(s, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.4),
            "Dans les deux premiers cas : jamais d'ecriture declenchee par le LLM — un port de lecture seule, systematiquement.",
            size=12, color=MUTED)
add_footer(s, 6, SECTION_NAMES[1])

# ============================================================ SLIDE 7 — SEPARATEUR PARTIE 2
section_divider(7, 2, "Architecture technique", "Hexagonale, stack, securite")

# ============================================================ SLIDE 8 — ARCHITECTURE HEXAGONALE
s = add_slide()
add_kicker(s, "Partie 2 — Architecture")
add_title(s, "Architecture hexagonale : ports & adaptateurs")
# Diagramme statique remplace par la video animee diagram_01_architecture (pulse voyageant
# Domain->Core->Infrastructure->Api) -- embed_videos.py la reconnait via le nom de cette forme.
# Carre (toutes les videos linkedin/ sont rendues en 1:1, frame_render.SIZE) : AddMediaObject2
# preserve l'aspect ratio natif de la video plutot que d'etirer vers des dimensions non
# carrees -- un slot rectangulaire large laisserait la video inseree minuscule et decalee.
slot_side = Inches(3.6)
slot_x = Inches(0.9) + (Inches(11.3) - slot_side) // 2
slot = add_box(s, slot_x, Inches(1.95), slot_side, slot_side, "▶  VIDEO : architecture hexagonale",
               fill=BG_ALT, line_color=BORDER, size=13, align=PP_ALIGN.CENTER, corner=True)
slot.name = "VIDEO_SLOT:diagram_01_architecture"
add_bullets(s, Inches(0.9), Inches(5.7), Inches(11.3), Inches(1.5), [
    "Le metier (Domain/Core) ne depend jamais de la technique — jamais l'inverse",
    "Domain ne sait meme pas qu'une base de donnees existe",
    "Remplacer SQL Server, Qdrant ou Ollama ne touche QUE Infrastructure — Core et Domain ne bougent pas",
    "Testabilite : la logique metier se teste sans infrastructure reelle (doubles Moq/FluentAssertions)",
], size=13, space_after=6)
add_footer(s, 8, SECTION_NAMES[2])

# ============================================================ SLIDE 9 — STACK TECHNIQUE
s = add_slide()
add_kicker(s, "Partie 2 — Architecture")
add_title(s, "Stack technique")
rows = [
    ("Backend", ".NET 8 · ASP.NET Core · EF Core"),
    ("Frontend", "Next.js 15 (BFF) · TailwindCSS · shadcn/ui"),
    ("Authentification", "JWT · ASP.NET Identity · cookie httpOnly"),
    ("Donnees relationnelles", "SQL Server"),
    ("Donnees vectorielles", "Qdrant (ANN/HNSW, cosinus)"),
    ("IA — embedding/reranking", "ONNX Runtime .NET pur (multilingue, cross-encodeur)"),
    ("IA — generation", "Ollama local (Router + Generator)"),
    ("Temps reel", "SSE (Server-Sent Events) — chat et notifications"),
    ("Deploiement", "Docker Compose (4 services)"),
]
tbl_shape = s.shapes.add_table(len(rows), 2, Inches(0.9), Inches(1.95), Inches(11.3), Inches(4.9))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.4)
tbl.columns[1].width = Inches(7.9)
for r, (k, v) in enumerate(rows):
    c0, c1 = tbl.cell(r, 0), tbl.cell(r, 1)
    c0.fill.solid(); c0.fill.fore_color.rgb = ACCENT_BG
    c1.fill.solid(); c1.fill.fore_color.rgb = BG_ALT
    for c, txt, bold, col in [(c0, k, True, PRIMARY), (c1, v, False, TEXT)]:
        c.margin_left = Pt(10); c.margin_top = Pt(4); c.margin_bottom = Pt(4)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; run = p.add_run(); run.text = txt
        run.font.size = Pt(13); run.font.bold = bold; run.font.color.rgb = col; run.font.name = FONT
tbl_shape.first_row = False
add_footer(s, 9, SECTION_NAMES[2])

# ============================================================ SLIDE 10 — SECURITE
s = add_slide()
add_kicker(s, "Partie 2 — Architecture")
add_title(s, "Securite : RBAC a portee + pattern BFF")
add_textbox(s, Inches(0.7), Inches(1.95), Inches(5.7), Inches(0.4), "RBAC + portee (DepartmentScopeGuard)", size=16, color=PRIMARY, bold=True)
code1 = ("public static bool CanAccessDepartment(...) =>\n"
         "    actor.Role switch\n"
         "    {\n"
         "        RoleType.QualityAdmin => true,\n"
         "        RoleType.HR => actor.DepartmentId\n"
         "            == targetDepartmentId,\n"
         "        _ => false\n"
         "    };")
box = add_box(s, Inches(0.7), Inches(2.4), Inches(5.7), Inches(2.35), code1, fill=RGBColor(0x0B,0x0B,0x0E),
              text_color=RGBColor(0xCB,0xE5,0xCB), size=11.5, align=PP_ALIGN.LEFT, bold=False, corner=True, line_color=BORDER)
for p in box.text_frame.paragraphs:
    for r in p.runs:
        r.font.name = "Consolas"
add_textbox(s, Inches(0.7), Inches(4.95), Inches(5.7), Inches(1.3),
            "Verifiee AVANT le role : un RH qui cible un dossier hors de son pole est refuse, quel que soit son role.",
            size=12.5, color=MUTED)

add_textbox(s, Inches(6.9), Inches(1.95), Inches(5.7), Inches(0.4), "Pattern BFF — le JWT ne quitte jamais le serveur", size=16, color=PRIMARY, bold=True)
add_box(s, Inches(6.9), Inches(2.55), Inches(1.55), Inches(0.85), "Navigateur", fill=BG_ALT, line_color=BORDER, size=11, corner=True)
add_arrow_right(s, Inches(8.45), Inches(2.975), length=Inches(0.32))
add_box(s, Inches(8.77), Inches(2.55), Inches(1.75), Inches(0.85), "Next.js\n(BFF)", fill=ACCENT_BG, line_color=PRIMARY, size=11, corner=True)
add_arrow_right(s, Inches(10.52), Inches(2.975), length=Inches(0.32))
add_box(s, Inches(10.84), Inches(2.55), Inches(1.56), Inches(0.85), "Arhia.Api", fill=BG_ALT, line_color=BORDER, size=11, corner=True)
add_textbox(s, Inches(6.9), Inches(3.65), Inches(5.7), Inches(0.35), "↳ appel serveur avec JWT + pose d'un cookie httpOnly", size=11.5, color=MUTED)
add_bullets(s, Inches(6.9), Inches(4.2), Inches(5.7), Inches(2.2), [
    "Le JWT n'est jamais accessible en JavaScript cote navigateur",
    "Meme un script XSS injecte ne peut pas le lire",
    "JWT valide : Issuer, Audience, Lifetime, cle de signature (les 4 verifications actives)",
], size=13, space_after=10)
add_footer(s, 10, SECTION_NAMES[2])

# ============================================================ SLIDE 11 — SEPARATEUR PARTIE 3
section_divider(11, 3, "Coeur IA — RAG & orchestration", "Le sous-systeme le plus dense du projet : 4 phases + 2 modeles + garde-fous")

# ============================================================ SLIDE 12 — RAG VUE D'ENSEMBLE
s = add_slide()
add_kicker(s, "Partie 3 — Coeur IA")
add_title(s, "Pipeline RAG — les 4 phases")
add_textbox(s, Inches(0.7), Inches(1.95), Inches(11.9), Inches(0.5),
            "RAG : au lieu de laisser le LLM repondre de memoire (donc halluciner), on cherche les vrais passages avant de generer.",
            size=13.5, color=MUTED)
# Diagramme statique (Chunking/Embedding/Recherche/Reranking) remplace par la video animee
# diagram_02_rag_pipeline (pulse traversant les 4 phases) -- couvre ce que les 2 anciennes
# slides "1/2"+"2/2" montraient separement en boites statiques.
# Carre (voir meme remarque qu'a la slide 8) : contraint par la hauteur disponible ici.
slot_side = Inches(2.5)
slot_x = Inches(0.7) + (Inches(11.9) - slot_side) // 2
slot = add_box(s, slot_x, Inches(2.65), slot_side, slot_side, "▶  VIDEO : pipeline RAG (4 phases)",
               fill=BG_ALT, line_color=BORDER, size=13, align=PP_ALIGN.CENTER, corner=True)
slot.name = "VIDEO_SLOT:diagram_02_rag_pipeline"

add_textbox(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(0.4), "Code reel — src/Arhia.Infrastructure/Rag/MarkdownChunker.cs, lignes 18-28", size=12, color=PRIMARY, bold=True)
code = ("public MarkdownChunker(Func<string,int> countTokens,\n"
        "    int maxTokensPerChunk = 400, double overlapRatio = 0.15)")
box = add_box(s, Inches(0.7), Inches(5.8), Inches(11.9), Inches(0.95), code, fill=RGBColor(0x0B,0x0B,0x0E),
              text_color=RGBColor(0xCB,0xE5,0xCB), size=13, align=PP_ALIGN.LEFT, bold=False, corner=True, line_color=BORDER)
for p in box.text_frame.paragraphs:
    for r in p.runs:
        r.font.name = "Consolas"
add_footer(s, 12, SECTION_NAMES[3])

# ============================================================ SLIDE 13 — RAG : POURQUOI 2 VITESSES
s = add_slide()
add_kicker(s, "Partie 3 — Coeur IA")
add_title(s, "Pipeline RAG — pourquoi deux vitesses (et un piege reel)")
# Les 2 boites "Phase 3/Phase 4" (redondantes avec la video de la slide 12) retirees ; le
# contenu propre a cette slide (comparaison + piege reel) remonte et s'agrandit pour occuper
# l'espace libere plutot que de laisser la moitie basse de la slide vide.
add_box(s, Inches(0.7), Inches(2.3), Inches(11.9), Inches(1.3), "Bi-encodeur (embedding) : rapide, separe, approximatif  →  Cross-encodeur (reranking) : lent, ensemble, precis",
        fill=BG_ALT, line_color=BORDER, size=18, corner=True, bold=False)
add_textbox(s, Inches(0.7), Inches(3.95), Inches(11.9), Inches(0.7),
            "Architecture a 2 etages : le bi-encodeur reduit vite tout le corpus a quelques candidats, "
            "le cross-encodeur les reordonne finement -- lent mais precis, applique seulement aux survivants.",
            size=16, color=MUTED, line_spacing=1.2)

add_textbox(s, Inches(0.7), Inches(5.0), Inches(11.9), Inches(0.4), "Piege reel retenu — src/Arhia.Infrastructure/Rag/OnnxRerankerAdapter.cs, ligne 62", size=13.5, color=WARN, bold=True)
add_textbox(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.4),
            "Un score de reranking eleve (jusqu'a 0.78 observe) ne garantit PAS que le chunk contient la reponse — "
            "seulement qu'il est thematiquement proche. Le score mesure une proximite, pas une verite : la decision finale "
            "\"source ou non\" relit le texte reellement genere, jamais le score seul.",
            size=15, color=MUTED, line_spacing=1.2)
add_footer(s, 13, SECTION_NAMES[3])

# ============================================================ SLIDE 14 — ORCHESTRATION
s = add_slide()
add_kicker(s, "Partie 3 — Coeur IA")
add_title(s, "Orchestration : Router + Generator")
add_textbox(s, Inches(0.7), Inches(1.95), Inches(11.9), Inches(0.4), "Deux roles distincts, deux appels LLM (Ollama, phi4-mini:3.8b)", size=13.5, color=MUTED)

add_box(s, Inches(0.7), Inches(2.5), Inches(5.6), Inches(2.5), "", fill=ACCENT_BG, line_color=PRIMARY, corner=True)
add_textbox(s, Inches(0.95), Inches(2.65), Inches(5.1), Inches(0.4), "Router — classification d'intention", size=15, color=TEXT, bold=True)
add_bullets(s, Inches(0.95), Inches(3.15), Inches(5.1), Inches(1.7), [
    "Sortie brute JAMAIS utilisee telle quelle",
    "Validee contre un enum ferme (3 valeurs)",
    "Tout le reste (vide, timeout, hallucination) → HORS_PERIMETRE par defaut",
], size=12.5, space_after=8)
add_badge(s, Inches(0.95), Inches(4.55), "FAIL-SAFE, PAS FAIL-OPEN", fill=GOOD, color=RGBColor(0x0B,0x0B,0x0E), w=Inches(3.4), size=11)

add_box(s, Inches(6.6), Inches(2.5), Inches(5.6), Inches(2.5), "", fill=ACCENT_BG, line_color=PRIMARY, corner=True)
add_textbox(s, Inches(6.85), Inches(2.65), Inches(5.1), Inches(0.4), "Generator — ecriture de la reponse", size=15, color=TEXT, bold=True)
add_bullets(s, Inches(6.85), Inches(3.15), Inches(5.1), Inches(1.7), [
    "Recoit uniquement le contexte deja filtre",
    "Streaming SSE — fragment par fragment",
    "Reponse toujours sourcee si RAG utilise",
], size=12.5, space_after=8)

add_textbox(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(1.2),
            "Limite mesuree et assumee : ~27% de mauvais classement (48 questions de test). Doubler les exemples few-shot "
            "n'a eu AUCUN effet mesurable (teste empiriquement, abandonne) — la capacite d'un modele 3,8B a suivre des "
            "regles explicites plafonne.", size=12.5, color=WARN, line_spacing=1.15)
add_footer(s, 14, SECTION_NAMES[3])

# ============================================================ SLIDE 15 — GARDE-FOU ANTI-HALLUCINATION
s = add_slide()
add_kicker(s, "Partie 3 — Coeur IA")
add_title(s, "Le garde-fou anti-hallucination — double porte de sortie")
add_box(s, Inches(0.9), Inches(2.05), Inches(2.6), Inches(0.8), "Recherche\nQdrant", fill=BG_ALT, line_color=BORDER, size=13, corner=True)
add_arrow_right(s, Inches(3.5), Inches(2.45), length=Inches(0.45))
add_box(s, Inches(4.05), Inches(1.85), Inches(3.1), Inches(1.2), "0 candidat trouve ?", fill=RGBColor(0x3A,0x1A,0x1A), line_color=WARN, size=13, corner=True)
add_arrow_right(s, Inches(7.25), Inches(2.45), length=Inches(0.45))
add_box(s, Inches(7.8), Inches(1.85), Inches(4.6), Inches(1.2), "→ refus immediat,\nGENERATOR JAMAIS APPELE", fill=RGBColor(0x1A,0x30,0x1E), line_color=GOOD, size=13, corner=True)

add_box(s, Inches(0.9), Inches(3.55), Inches(2.6), Inches(0.8), "Reranking\n(score)", fill=BG_ALT, line_color=BORDER, size=13, corner=True)
add_arrow_right(s, Inches(3.5), Inches(3.95), length=Inches(0.45))
add_box(s, Inches(4.05), Inches(3.35), Inches(3.1), Inches(1.2), "Rien ne passe\nle seuil (0.01) ?", fill=RGBColor(0x3A,0x1A,0x1A), line_color=WARN, size=13, corner=True)
add_arrow_right(s, Inches(7.25), Inches(3.95), length=Inches(0.45))
add_box(s, Inches(7.8), Inches(3.35), Inches(4.6), Inches(1.2), "→ refus immediat,\nGENERATOR JAMAIS APPELE", fill=RGBColor(0x1A,0x30,0x1E), line_color=GOOD, size=13, corner=True)

add_textbox(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(0.4), "Code reel — src/Arhia.Core/UseCases/AnswerConversationUseCase.cs, lignes 179-194", size=12, color=PRIMARY, bold=True)
code = ("if (candidates.Count == 0)\n"
        "    return new DocumentaryPreparation(false, null, ...);\n"
        "// ... reranking, filtrage par MinimumRelevanceThreshold = 0.01f ...\n"
        "if (best.Count == 0)\n"
        "    return new DocumentaryPreparation(false, null, ...);")
box = add_box(s, Inches(0.7), Inches(5.48), Inches(11.9), Inches(1.35), code, fill=RGBColor(0x0B,0x0B,0x0E),
              text_color=RGBColor(0xCB,0xE5,0xCB), size=12.5, align=PP_ALIGN.LEFT, bold=False, corner=True, line_color=BORDER)
for p in box.text_frame.paragraphs:
    for r in p.runs:
        r.font.name = "Consolas"
add_footer(s, 15, SECTION_NAMES[3])

# ============================================================ SLIDE 16 — SEPARATEUR PARTIE 4
section_divider(16, 4, "Resultats & deploiement", "Ce qui est verifie en conditions reelles, et ce qui reste ouvert")

# ============================================================ SLIDE 17 — RESULTATS
s = add_slide()
add_kicker(s, "Partie 4 — Resultats")
add_title(s, "Resultats mesures — verifies, pas juste supposes")
stats = [("244 / 245", "tests automatises verts", GOOD), ("0", "warning au build", GOOD), ("4", "phases RAG toutes obligatoires", PRIMARY), ("~27%", "erreurs de routage — limite connue et assumee", WARN)]
x = Inches(0.7); w = Inches(2.85)
for i, (num, label, col) in enumerate(stats):
    l = x + i * (w + Inches(0.15))
    add_box(s, l, Inches(2.0), w, Inches(1.5), "", fill=ACCENT_BG, line_color=col, corner=True)
    add_textbox(s, l + Inches(0.15), Inches(2.15), w - Inches(0.3), Inches(0.7), num, size=28, color=col, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(s, l + Inches(0.15), Inches(2.85), w - Inches(0.3), Inches(0.6), label, size=11, color=MUTED, align=PP_ALIGN.CENTER)
add_bullets(s, Inches(0.7), Inches(3.95), Inches(11.9), Inches(2.9), [
    "Chat en streaming reel verifie en HTTP de bout en bout : question -> reponse sourcee, token par token",
    "Anti-hallucination code (pas seulement prompt) : verifie par des tests unitaires dedies",
    ("Limite du routeur documentee comme un compromis assume, pas cachee — mode de defaillance \"gracieusement faux\"", True),
    "Piste testee et abandonnee : plus d'exemples few-shot dans le prompt -> aucun effet mesurable",
    "Base de dev peuplee avec un scenario realiste (5 poles, 25 collaborateurs, 2 templates reels) — RBAC et portee departement verifies en HTTP reel sur chaque creation",
    ("Le seul echec (1/245) est un flake pre-existant et documente (RagPipelineIntegrationTests, accumulation Qdrant persistante), sans rapport avec le code recent", True),
], size=13, space_after=11)
add_footer(s, 17, SECTION_NAMES[4])

# ============================================================ SLIDE 18 — DEPLOIEMENT
s = add_slide()
add_kicker(s, "Partie 4 — Resultats")
add_title(s, "Deploiement : une seule commande, verifiee sur base fraiche")
services = [("sqlserver", "1433", "Entites metier"), ("qdrant", "6333", "Index vectoriel RAG"), ("api", "5080", "Backend .NET"), ("frontend", "3000", "Next.js (BFF)")]
x = Inches(0.7); w = Inches(2.85)
for i, (name, port, role) in enumerate(services):
    l = x + i * (w + Inches(0.15))
    add_box(s, l, Inches(2.05), w, Inches(1.35), f"{name}\n:{port}", fill=ACCENT_BG, line_color=PRIMARY, size=15, corner=True)
    add_textbox(s, l + Inches(0.1), Inches(3.5), w - Inches(0.2), Inches(0.5), role, size=11, color=MUTED, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.7), Inches(4.25), Inches(11.9), Inches(0.4), "+ Ollama natif sur l'hote (jamais conteneurise) — rejoint via host.docker.internal", size=13, color=MUTED)
add_bullets(s, Inches(0.7), Inches(4.85), Inches(11.9), Inches(2.0), [
    "docker compose up -d --build : demarre toute la pile en une commande",
    "Migrations EF Core auto-appliquees au demarrage (idempotent)",
    "Verifie en conditions reelles completes : inscription -> connexion -> chat sourcee, a travers toute la pile conteneurisee",
], size=13.5, space_after=10)
add_footer(s, 18, SECTION_NAMES[4])

# ============================================================ SLIDE 19 — BILAN & PERSPECTIVES
s = add_slide()
add_kicker(s, "Partie 4 — Resultats")
add_title(s, "Bilan & perspectives")
add_textbox(s, Inches(0.7), Inches(2.0), Inches(5.6), Inches(0.4), "Livre et verifie", size=16, color=GOOD, bold=True)
add_bullets(s, Inches(0.7), Inches(2.5), Inches(5.6), Inches(4.5), [
    "Socle metier complet (RBAC, workflows, circuit qualite)",
    "Pipeline RAG 4 phases, ONNX de bout en bout",
    "Orchestration + garde-fous anti-hallucination",
    "Interface complete par role — dashboard, workflows, employes (au-dela du seul chat)",
    "Administration : relance de l'ingestion RAG depuis l'UI (QualityAdmin)",
    "Frontend temps reel (chat + notifications SSE) + mode sombre",
    "Deploiement Docker Compose complet",
    "Base de dev peuplee et verifiee (scenario realiste, RBAC teste en conditions reelles)",
], size=13.5, space_after=12, bullet_color=GOOD)
add_textbox(s, Inches(6.9), Inches(2.0), Inches(5.6), Inches(0.4), "Prochaines etapes", size=16, color=PRIMARY, bold=True)
add_bullets(s, Inches(6.9), Inches(2.5), Inches(5.6), Inches(4.5), [
    "Ameliorer la precision du routeur (nouvelle approche a evaluer)",
    "Templates, admin complementaire, polish chat/notifications (phases frontend restantes)",
    "3 cas particuliers metier (mutation, suspension, pole vacant)",
    "Profil Ollama entreprise (modeles plus capables)",
], size=13.5, space_after=12)
add_footer(s, 19, SECTION_NAMES[4])

# ============================================================ SLIDE 20 — MERCI
s = add_slide()
add_logo(s, ARHIA_LOGO, Inches(5.07), Inches(2.1), Inches(0.9))
add_textbox(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(1.0), "Merci de votre attention", size=40, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
add_textbox(s, Inches(0.9), Inches(4.3), Inches(11.5), Inches(0.5), "Questions & demonstration en direct", size=16, color=MUTED, align=PP_ALIGN.CENTER)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.17), Inches(4.95), Inches(1.0), Pt(3))
line.fill.solid(); line.fill.fore_color.rgb = PRIMARY; line.line.fill.background(); line.shadow.inherit = False

sortie = os.path.join(os.path.dirname(os.path.abspath(__file__)), "arhia_Soutenance.pptx")
prs.save(sortie)
print("OK -", len(prs.slides.__iter__.__self__._sldIdLst), "slides ->", sortie)
