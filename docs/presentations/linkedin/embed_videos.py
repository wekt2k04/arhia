# -*- coding: utf-8 -*-
"""
Insere des videos (autoplay + boucle continue, sans barre de controle
persistante) dans des slides cibles d'un .pptx existant.

Cible une slide par deux mecanismes, au choix (au moins un requis) :
  1. Mot-cle dans les notes de la slide : une ligne contenant
     [[video:NOM]] (NOM sans extension) declenche l'insertion de
     NOM.mp4 depuis le dossier fourni.
  2. Forme "espace reservé" : une forme dont le nom commence par
     VIDEO_SLOT: (ex. "VIDEO_SLOT:diagram_01_architecture") -- sa
     position/taille sert de cadre pour la video, puis elle est
     supprimee (remplacee par la video). A defaut de forme, une
     position/taille par defaut centree est utilisee.

Mecanique choisie -- COM (pywin32) plutot que python-pptx pur pour l'ecriture :
python-pptx n'expose pas les reglages de lecture (autoplay/boucle) via son
API haut niveau, et un bug documente (scanny/python-pptx#954 -- <p:timing>
duplique) peut corrompre le fichier si on les injecte a la main en XML brut.
PowerPoint lui-meme, pilote via Shape.AnimationSettings.PlaySettings,
genere une XML garantie valide -- verifie ici par un aller-retour reel
(insertion -> sauvegarde -> fermeture -> reouverture -> relecture des
proprietes, toutes confirmees identiques).

Limite connue, assumee honnetement plutot que passee sous silence : au-dela
d'autoplay + boucle, il n'existe pas de reglage COM distinct "masquer les
controles" pour une video en lecture automatique en boucle -- en mode
Diaporama, PowerPoint n'affiche de barre de lecture qu'au survol de la
souris (jamais une barre persistante), ce qui est deja le comportement
obtenu avec ce script sans reglage supplementaire.

Usage :
    python embed_videos.py <entree.pptx> <dossier_videos> [--output <sortie.pptx>]

Necessite pywin32 (deja present sur ce poste) et PowerPoint installe.
"""
import argparse
import os
import re
import sys

from pptx import Presentation

VIDEO_NOTE_RE = re.compile(r"\[\[video:([\w-]+)\]\]")
SLOT_PREFIX = "VIDEO_SLOT:"

# Shapes.AddMediaObject2(Left/Top/Width/Height) attend des POINTS (comme Shape.Left/.Top/
# .Width/.Height eux-memes releus ensuite), PAS des pixels a 96 dpi -- confirme empiriquement
# (position/taille observees apres insertion = exactement les valeurs "px" calculees avec
# l'ancienne constante 9525, reinterpretees comme des points : le bug etait bien la, pas
# ailleurs). 1 point = 1/72 pouce = 12700 EMU.
EMU_PER_POINT = 12700


def find_targets(pptx_path, videos_dir):
    """Phase 1 (lecture seule, python-pptx) : determine quelle video va sur
    quelle slide, et a quelle position/taille (EMU). Ne touche jamais au
    fichier -- l'ecriture se fait entierement en phase 2 (COM)."""
    prs = Presentation(pptx_path)
    targets = []
    for idx, slide in enumerate(prs.slides):
        video_name = None

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            m = VIDEO_NOTE_RE.search(notes)
            if m:
                video_name = m.group(1)

        slot_shape = None
        for shape in slide.shapes:
            if shape.name.startswith(SLOT_PREFIX):
                slot_shape = shape
                if video_name is None:
                    video_name = shape.name[len(SLOT_PREFIX):]
                break

        if video_name is None:
            continue

        video_path = os.path.join(videos_dir, f"{video_name}.mp4")
        if not os.path.isfile(video_path):
            print(f"  ATTENTION slide {idx + 1} : {video_name}.mp4 introuvable dans {videos_dir}, ignoree.")
            continue

        if slot_shape is not None:
            left, top, width, height = slot_shape.left, slot_shape.top, slot_shape.width, slot_shape.height
        else:
            # Position par defaut : cadre centre, ~40% de la largeur de la slide.
            sw, sh = prs.slide_width, prs.slide_height
            width = int(sw * 0.4)
            height = int(width)  # nos videos sont carrees (1:1)
            left = (sw - width) // 2
            top = (sh - height) // 2

        targets.append({
            "slide_index_1based": idx + 1,
            "video_path": os.path.abspath(video_path),
            "slot_shape_name": slot_shape.name if slot_shape is not None else None,
            "left": left, "top": top, "width": width, "height": height,
        })

    return targets


def embed_via_com(pptx_path, output_path, targets):
    """Phase 2 (ecriture, COM) : PlaySettings genere une XML de timing
    garantie valide par PowerPoint lui-meme -- pas de manipulation XML
    manuelle ici. Ouvre une seule fois, traite toutes les slides, sauvegarde
    une seule fois."""
    import win32com.client

    app = win32com.client.Dispatch("PowerPoint.Application")
    app.Visible = True  # PlaySettings/AddMediaObject2 se sont montres peu fiables invisibles
    pres = None
    try:
        # ReadOnly=False : necessaire meme si output_path == pptx_path n'est pas le cas le plus
        # frequent -- PowerPoint refuse un SaveAs vers le MEME nom qu'un fichier ouvert ReadOnly
        # ("must be saved with a different name"), rencontre en pratique en ecrasant le pptx de
        # soutenance sur place. Sans consequence pour le cas output != pptx_path (fichier source
        # jamais modifie sur disque avant le SaveAs explicite vers output_path de toute facon).
        pres = app.Presentations.Open(pptx_path, False, False, False)
        for t in targets:
            slide = pres.Slides.Item(t["slide_index_1based"])

            if t["slot_shape_name"] is not None:
                for shape in slide.Shapes:
                    if shape.Name == t["slot_shape_name"]:
                        shape.Delete()
                        break

            left_pt = t["left"] / EMU_PER_POINT
            top_pt = t["top"] / EMU_PER_POINT
            width_pt = t["width"] / EMU_PER_POINT
            height_pt = t["height"] / EMU_PER_POINT

            shape = slide.Shapes.AddMediaObject2(
                t["video_path"], False, True, left_pt, top_pt, width_pt, height_pt)
            shape.AnimationSettings.PlaySettings.PlayOnEntry = True
            shape.AnimationSettings.PlaySettings.LoopUntilStopped = True
            shape.AnimationSettings.PlaySettings.HideWhileNotPlaying = False
            print(f"  Slide {t['slide_index_1based']} : {os.path.basename(t['video_path'])} inseree "
                  f"({'slot ' + t['slot_shape_name'] if t['slot_shape_name'] else 'position par defaut'}).")

        pres.SaveAs(output_path)
        print(f"OK -> {output_path}")
    finally:
        if pres:
            pres.Close()
        app.Quit()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx", help="Fichier .pptx d'entree (jamais modifie directement)")
    ap.add_argument("videos_dir", help="Dossier contenant les .mp4 a inserer")
    ap.add_argument("--output", help="Fichier .pptx de sortie (defaut : <entree>_with_videos.pptx)")
    args = ap.parse_args()

    # Chemins absolus obligatoires : le process PowerPoint lance via COM a son propre
    # repertoire de travail, un chemin relatif y est resolu ailleurs -- deja rencontre
    # avec export_video.ps1 (piege documente dans .claude/HANDOFF/NEXT_SESSION.md).
    pptx_path = os.path.abspath(args.pptx)
    videos_dir = os.path.abspath(args.videos_dir)
    output = os.path.abspath(args.output or re.sub(r"\.pptx$", "_with_videos.pptx", pptx_path))

    print(f"Analyse de {pptx_path}...")
    targets = find_targets(pptx_path, videos_dir)
    if not targets:
        print("Aucune slide cible trouvee ([[video:NOM]] dans les notes, ou forme nommee "
              f"'{SLOT_PREFIX}NOM'). Rien a faire.")
        return 1

    print(f"{len(targets)} slide(s) cible(s) trouvee(s). Insertion via PowerPoint (COM)...")
    embed_via_com(pptx_path, output, targets)
    return 0


if __name__ == "__main__":
    sys.exit(main())
